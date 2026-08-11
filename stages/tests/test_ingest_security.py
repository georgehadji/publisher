"""
U5/S9 -- DOCX ingest hardening (docs/ARCHITECTURE_UPLIFT_PLAN.md §3 U5 S9).

The upload route caps the wire bytes; these tests cap what a stage is allowed
to TRUST once the bytes are on disk. A hostile DOCX is a ZIP archive: it can
declare huge uncompressed sizes or a huge entry count while staying small on
the wire, and python-docx will decompress whatever the central directory
promises. `ingest` refuses both before lxml ever parses.
"""

from __future__ import annotations

import io
import json
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import pytest

from publisher_stages import ErrorKind, StageCtx, StageError
from stages.ingest_stage import ingest


def _ctx(tmp_path: Path) -> StageCtx:
    return StageCtx(
        build_id="test-ingest-sec",
        cache_key="test",
        deadline=datetime.now(timezone.utc),
        memory_budget_mb=128,
        work_dir=str(tmp_path),
        cas_root=str(tmp_path / "cas"),
        allow_stub_engines=False,
    )


def _real_docx() -> bytes:
    """A minimal valid DOCX (python-docx), as the upload route would store it."""
    import docx

    document = docx.Document()
    document.add_heading("SECURITY TEST", level=1)
    document.add_paragraph("Body text for the ingest security test.")
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _write_docx(tmp_path: Path, data: bytes) -> str:
    path = tmp_path / "m.docx"
    path.write_bytes(data)
    return str(path)


def test_ingest_accepts_a_normal_docx(tmp_path):
    """Control: a legitimate DOCX passes the zip-bomb guard and ingests."""
    cas_root = tmp_path / "cas"
    cas_root.mkdir()
    path = _write_docx(tmp_path, _real_docx())
    result = ingest(_ctx(tmp_path), docx_path=path)
    assert len(result.artifacts) == 1
    assert result.artifacts[0].kind == "source"
    assert result.artifacts[0].media_type == "application/json"


def test_ingest_rejects_many_entries_zip_bomb(tmp_path):
    """S9: an archive with more than MAX_DOCX_ENTRIES entries is refused."""
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as zf:
        for i in range(5000):  # default cap is 4096
            zf.writestr(f"word/part{i}.xml", "<w:p/>")
    path = _write_docx(tmp_path, bomb.getvalue())

    with pytest.raises(StageError) as exc:
        ingest(_ctx(tmp_path), docx_path=path)
    assert exc.value.kind == ErrorKind.BAD_INPUT
    assert "zip bomb" in exc.value.message


def test_ingest_rejects_huge_decompressed_size(tmp_path):
    """S9: an archive declaring more uncompressed bytes than the cap is refused."""
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as zf:
        # 300 MB of zeros -- ~300 KB on the wire, over the 256 MB cap on disk.
        zf.writestr("word/document.xml", b"\x00" * (300 * 1024 * 1024))
    path = _write_docx(tmp_path, bomb.getvalue())

    with pytest.raises(StageError) as exc:
        ingest(_ctx(tmp_path), docx_path=path)
    assert exc.value.kind == ErrorKind.BAD_INPUT
    assert "zip bomb" in exc.value.message


def test_ingest_rejects_corrupt_archive(tmp_path):
    """S9: bytes that are not a ZIP at all are named as such, not left to blow
    up python-docx three stack frames later."""
    path = _write_docx(tmp_path, b"this is not a zip archive")
    with pytest.raises(StageError) as exc:
        ingest(_ctx(tmp_path), docx_path=path)
    assert exc.value.kind == ErrorKind.BAD_INPUT
    assert "ZIP/DOCX" in exc.value.message
