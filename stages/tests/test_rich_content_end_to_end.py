"""
A manuscript with a picture, a table and a footnote, ingest to PDF.

This is the test the pipeline did not have. Every stage passed its own checks
while images and footnotes were being dropped, because the one hard gate that
runs on every build compares *text*, and:

  * a picture has no text to miss;
  * a footnote body lives in `word/footnotes.xml`, which never reached the AST,
    so neither side of the gate ever saw it;
  * table cells did reach the AST -- as loose paragraphs -- so their text
    matched and the lost table structure went unnoticed.

The assertions are therefore deliberately not "the stage returned an artifact".
They are: the integrity gate passes on rich content, the PDF has an embedded
image, and the footnote and table text are on the rendered page.
"""

from __future__ import annotations

import json
import shutil
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import docx
import pytest

from publisher_stages import StageCtx

HAVE_PANDOC = bool(shutil.which("pandoc"))
HAVE_TYPST = bool(shutil.which("typst"))

PNG_1PX = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08"
    b"\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00"
    b"\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
)

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
FOOTNOTE_TEXT = "The archive copy differs in three places."
FOOTNOTES_XML = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:footnotes xmlns:w="{W_NS}">
  <w:footnote w:type="separator" w:id="0"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>
  <w:footnote w:id="2"><w:p><w:r><w:t>{FOOTNOTE_TEXT}</w:t></w:r></w:p></w:footnote>
</w:footnotes>""".encode("utf-8")

PROSE = ("This chapter runs long enough to read as body text rather than front "
         "matter, which the structure heuristic decides by paragraph length. ") * 3


def _rich_docx(tmp_path: Path) -> Path:
    document = docx.Document()
    document.add_paragraph("CHAPTER ONE")

    para = document.add_paragraph()
    para.add_run(PROSE)
    para.add_run("emphatic").italic = True

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Edition"
    table.cell(0, 1).text = "Copies"
    table.cell(1, 0).text = "Aldine"
    table.cell(1, 1).text = "1200"

    plate = tmp_path / "plate.png"
    plate.write_bytes(PNG_1PX)
    document.add_picture(str(plate))

    plain = tmp_path / "plain.docx"
    document.save(str(plain))

    # python-docx cannot author footnotes; attach one to the prose paragraph.
    with zipfile.ZipFile(plain) as zin:
        parts = {n: zin.read(n) for n in zin.namelist()}
    body = parts["word/document.xml"].decode("utf-8")
    close = body.index("</w:r>", body.index(">emphatic<")) + len("</w:r>")
    body = body[:close] + '<w:r><w:footnoteReference w:id="2"/></w:r>' + body[close:]
    parts["word/document.xml"] = body.encode("utf-8")
    parts["word/footnotes.xml"] = FOOTNOTES_XML
    parts["word/_rels/document.xml.rels"] = (
        parts["word/_rels/document.xml.rels"].decode("utf-8").replace(
            "</Relationships>",
            '<Relationship Id="rIdFn" Type="http://schemas.openxmlformats.org/'
            'officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/>'
            "</Relationships>")
    ).encode("utf-8")
    parts["[Content_Types].xml"] = (
        parts["[Content_Types].xml"].decode("utf-8").replace(
            "</Types>",
            '<Override PartName="/word/footnotes.xml" ContentType="application/'
            'vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>'
            "</Types>")
    ).encode("utf-8")

    rich = tmp_path / "rich.docx"
    with zipfile.ZipFile(rich, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in parts.items():
            zout.writestr(name, data)
    return rich


def _ctx(tmp_path: Path) -> StageCtx:
    return StageCtx(
        build_id="rich",
        cache_key="rich",
        deadline=datetime.now(timezone.utc),
        memory_budget_mb=512,
        work_dir=str(tmp_path / "work"),
        cas_root=str(tmp_path / "cas"),
    )


def _cas_path(ctx: StageCtx, digest: str) -> Path:
    return Path(ctx.cas_root) / digest[:2] / digest[2:4] / digest


def _ingest_to_html(tmp_path: Path) -> tuple[StageCtx, dict, str]:
    """Run ingest -> extract, returning (ctx, source AST, HTML)."""
    from stages.ingest_stage import ingest
    from stages.extract_stage import extract

    ctx = _ctx(tmp_path)
    Path(ctx.work_dir).mkdir(parents=True, exist_ok=True)

    ingested = ingest(ctx, docx_path=str(_rich_docx(tmp_path)))
    source_path = _cas_path(ctx, ingested.artifacts[0].hash)
    source = json.loads(source_path.read_bytes())

    extracted = extract(ctx, source=str(source_path))
    html = _cas_path(ctx, extracted.artifacts[0].hash).read_text(encoding="utf-8")
    return ctx, source, html


def test_rich_content_passes_the_text_integrity_gate(tmp_path):
    """The gate compares HTML text against AST text. Footnote bodies and table
    cells are now in the AST, so extract has to render both -- an unhandled node
    type drops its text silently and the gate fails the build."""
    from stages.structure_stage import ast_assemble

    ctx, source, html = _ingest_to_html(tmp_path)
    work = Path(ctx.work_dir)
    (work / "source.json").write_text(json.dumps(source), encoding="utf-8")
    (work / "doc.html").write_text(html, encoding="utf-8")

    # Raises StageError(ENGINE_BUG) on any mismatch -- no flag, no override.
    result = ast_assemble(ctx, html=str(work / "doc.html"),
                          source=str(work / "source.json"))
    assert result.metrics["integrity_ok"] == 1.0
    assert FOOTNOTE_TEXT in html
    assert "Aldine" in html


def test_extract_emits_a_real_img_and_a_real_table(tmp_path):
    """A figure used to render as `<figure><figcaption/></figure>` -- an empty
    box. Nothing failed, because a picture contributes no text to compare."""
    _, _, html = _ingest_to_html(tmp_path)

    assert '<img src="media/' in html
    assert "<table>" in html and "<tr>" in html
    assert '<span class="footnote">' in html
    assert "<em>emphatic</em>" in html


def test_media_materializes_out_of_cas(tmp_path):
    """The HTML references images by hash because the work directory it was
    produced in is scratch. Renderers pull the bytes back out."""
    from stages.media import materialize_media

    ctx, _, html = _ingest_to_html(tmp_path)
    out = Path(ctx.work_dir) / "render"
    written = materialize_media(html, ctx.cas_root, out)

    assert len(written) == 1
    assert written[0].read_bytes() == PNG_1PX
    assert written[0].parent == out / "media"


def test_a_referenced_figure_missing_from_cas_is_an_error(tmp_path):
    """Guards the guard: a blank rectangle where a plate should be looks like a
    finished book, so a missing blob has to stop the build."""
    from stages.media import materialize_media

    html = '<img src="media/' + "0" * 64 + '.png"/>'
    with pytest.raises(FileNotFoundError, match="not in CAS"):
        materialize_media(html, tmp_path / "empty-cas", tmp_path / "out")


@pytest.mark.skipif(not (HAVE_PANDOC and HAVE_TYPST),
                    reason="needs pandoc and typst")
def test_typst_renders_image_footnote_and_table_into_the_pdf(tmp_path):
    """End of the line: the things that were being dropped are in the PDF.

    Read back out of the rendered file rather than asserted on the Typst source,
    because "pandoc emitted #footnote" and "the note is on the page" are
    different claims, and only the second one is the deliverable.
    """
    import fitz
    from stages.resolve_stage import resolve
    from stages.structure_stage import ast_assemble
    from stages.typst_stages import paginate_typst

    ctx, source, html = _ingest_to_html(tmp_path)
    work = Path(ctx.work_dir)
    (work / "source.json").write_text(json.dumps(source), encoding="utf-8")
    (work / "doc.html").write_text(html, encoding="utf-8")

    # The real chain: the render path consumes `doc-effective/1`, which is only
    # reachable through the integrity gate. Rendering straight off the source
    # AST would test a path no build can take.
    assembled = ast_assemble(ctx, html=str(work / "doc.html"),
                             source=str(work / "source.json"))
    ast_path = _cas_path(ctx, next(a for a in assembled.artifacts
                                   if a.kind == "ast").hash)
    resolved = resolve(ctx, ast=str(ast_path))
    doc_path = _cas_path(ctx, resolved.artifacts[0].hash)

    result = paginate_typst(ctx, doc_path=str(doc_path))
    pdf = _cas_path(ctx, next(a for a in result.artifacts if a.kind == "pdf").hash)

    with fitz.open(pdf) as rendered:
        text = "".join(page.get_text() for page in rendered)
        images = [img for page in rendered for img in page.get_images(full=True)]

    assert images, "no embedded image in the PDF -- the figure was dropped"
    assert FOOTNOTE_TEXT in text, "footnote body is not on the page"
    assert "Aldine" in text and "1200" in text, "table content is missing"
    # The call is set by Typst, not spliced into the prose by the extractor.
    assert "emphatic1" not in text
