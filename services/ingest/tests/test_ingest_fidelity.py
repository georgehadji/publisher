"""
W9: what ingest used to lose while every check stayed green.

Each case is content a real manuscript carries: a symbol-font arrow (Word's
AutoCorrect for "-->"), headings bold only through their style, a caption typed
under its picture, and drawn lines and arrows the AST has no node for.
"""

from __future__ import annotations

import json
import zipfile
from pathlib import Path

import docx
import pytest
from docx.enum.style import WD_STYLE_TYPE
from lxml import etree

from publisher_ingest.docx_rich import W, sha256_hex, source_texts, sym_char
from publisher_ingest.docx_to_ast import IngestError, docx_to_ast, drawings_dropped

PROSE = "Filler prose to clear the substantive threshold. " * 8
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
V_NS = "urn:schemas-microsoft-com:vml"
PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08"
    b"\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00"
    b"\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
)


def _run_xml(inner: str) -> etree._Element:
    return etree.fromstring(f'<w:r xmlns:w="{W_NS}" xmlns:v="{V_NS}">{inner}</w:r>')


def _save(document, tmp_path: Path) -> Path:
    path = tmp_path / "manuscript.docx"
    document.save(str(path))
    return path


def _texts(ast) -> str:
    return json.dumps(ast, ensure_ascii=False)


# ── w:sym ────────────────────────────────────────────────────────────────

def _with_symbol(tmp_path: Path, font: str, code: str) -> Path:
    document = docx.Document()
    document.add_paragraph("CHAPTER ONE")
    para = document.add_paragraph("Before ")
    para._p.append(_run_xml(f'<w:sym w:font="{font}" w:char="{code}"/>'))
    para.add_run(" after. " + PROSE)
    return _save(document, tmp_path)


@pytest.mark.parametrize("font, code, char", [
    ("Wingdings", "F0E0", "→"),   # Word's AutoCorrect for "-->"
    ("Wingdings", "00E0", "→"),   # the same arrow, addressed without the F0 page
    ("Symbol", "F061", "α"),
    ("Symbol", "F0B3", "≥"),
    ("Times New Roman", "2014", "—"),   # an ordinary font: the code is Unicode
])
def test_a_symbol_reaches_the_ast_and_the_oracle(tmp_path, font, code, char):
    path = _with_symbol(tmp_path, font, code)

    assert f"Before {char} after." in _texts(docx_to_ast(path))
    # The oracle counts it too, so a walk that dropped it would fail ingest.
    assert any(f"Before {char} after." in t for t in source_texts(docx.Document(str(path))))


def test_an_unmapped_symbol_is_refused_by_name(tmp_path):
    with pytest.raises(IngestError, match="Wingdings symbol .code F04A"):
        docx_to_ast(_with_symbol(tmp_path, "Wingdings", "F04A"))


def test_sym_char_reads_the_low_byte_of_a_symbol_font():
    sym = etree.fromstring(f'<w:sym xmlns:w="{W_NS}" w:font="Symbol" w:char="F057"/>')
    assert sym.tag == f"{W}sym"
    assert sym_char(sym) == "Ω"


# ── bold through a style ────────────────────────────────────────────────

def test_headings_bold_only_through_their_style_open_chapters(tmp_path):
    document = docx.Document()
    styles = document.styles
    base = styles.add_style("Section Base", WD_STYLE_TYPE.PARAGRAPH)
    base.base_style = styles["Normal"]
    base.font.bold = True
    section = styles.add_style("Section", WD_STYLE_TYPE.PARAGRAPH)
    section.base_style = base          # bold two steps up the chain
    strong = styles.add_style("Strong Words", WD_STYLE_TYPE.CHARACTER)
    strong.font.bold = True

    document.add_paragraph("1. Εισαγωγή", style="Section")
    document.add_paragraph(PROSE)
    heading = document.add_paragraph()
    heading.add_run("2. Second part").style = strong   # bold by character style
    document.add_paragraph(PROSE)
    plain = document.add_paragraph(style="Section")
    plain.add_run("3. Not a heading").bold = False      # direct formatting wins
    document.add_paragraph(PROSE)

    titles = [c["attrs"]["title"] for c in docx_to_ast(_save(document, tmp_path))["body"]]

    assert titles == ["1. Εισαγωγή", "2. Second part"]


# ── a caption typed under its picture ───────────────────────────────────

def test_a_typed_caption_moves_into_its_figure(tmp_path):
    image = tmp_path / "plate.png"
    image.write_bytes(PNG)
    document = docx.Document()
    document.add_paragraph("CHAPTER ONE")
    document.add_paragraph(PROSE)
    document.add_picture(str(image))
    document.add_paragraph("Εικόνα 6: Ο χάρτης της περιοχής")
    document.add_picture(str(image))
    document.add_paragraph("The map shows the region.")   # prose, not a caption

    ast = docx_to_ast(_save(document, tmp_path), store_media=lambda b, t, n: sha256_hex(b))
    content = ast["body"][0]["content"]

    figures = [n for n in content if n["type"] == "figure"]
    assert figures[0]["attrs"]["caption"] == "Εικόνα 6: Ο χάρτης της περιοχής"
    assert "caption" not in figures[1]["attrs"]
    # Moved, not copied: no paragraph still holds it.
    assert _texts(content).count("Εικόνα 6") == 1
    assert content[-1]["type"] == "paragraph"


# ── drawings the AST cannot hold ────────────────────────────────────────

def test_dropped_lines_are_counted_and_text_boxes_are_not(tmp_path):
    document = docx.Document()
    document.add_paragraph("CHAPTER ONE")
    para = document.add_paragraph(PROSE)
    para._p.append(_run_xml('<w:pict><v:line from="0,0" to="10,10"/></w:pict>'))
    para._p.append(_run_xml('<w:pict><v:shape><v:textbox><w:txbxContent>'
                            '<w:p><w:r><w:t>Boxed words</w:t></w:r></w:p>'
                            '</w:txbxContent></v:textbox></v:shape></w:pict>'))
    path = _save(document, tmp_path)

    assert drawings_dropped(path) == {"shapes": 1, "pictures": 0, "objects": 0}
    assert "Boxed words" in _texts(docx_to_ast(path))


def test_the_ingest_stage_warns_about_dropped_drawings(tmp_path):
    from datetime import datetime, timezone
    from publisher_stages import StageCtx
    from stages.ingest_stage import ingest

    document = docx.Document()
    document.add_paragraph("CHAPTER ONE")
    document.add_paragraph(PROSE)._p.append(
        _run_xml('<w:pict><v:line from="0,0" to="10,10"/></w:pict>'))
    path = _save(document, tmp_path)
    ctx = StageCtx(build_id="t", deterministic_seed="t", deadline=datetime.now(timezone.utc),
                   memory_budget_mb=512, work_dir=str(tmp_path / "work"),
                   cas_root=str(tmp_path / "cas"))

    result = ingest(ctx, docx_path=str(path))

    assert result.metrics["shapes_dropped"] == 1.0
    assert [w.code for w in result.warnings] == ["drawings-dropped"]


# ── endnotes ────────────────────────────────────────────────────────────

ENDNOTES_XML = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:endnotes xmlns:w="{W_NS}">
  <w:endnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:endnote>
  <w:endnote w:id="1"><w:p><w:r><w:t>An endnote's words.</w:t></w:r></w:p></w:endnote>
</w:endnotes>""".encode("utf-8")


def test_a_word_endnote_becomes_an_endnote(tmp_path):
    """It used to become a footnote: printed at the page foot, not the chapter end."""
    document = docx.Document()
    document.add_paragraph("CHAPTER ONE")
    document.add_paragraph("Cited here." + PROSE)._p.append(
        _run_xml('<w:endnoteReference w:id="1"/>'))
    src = _save(document, tmp_path)

    with zipfile.ZipFile(src) as zin:
        parts = {n: zin.read(n) for n in zin.namelist()}
    parts["word/endnotes.xml"] = ENDNOTES_XML
    parts["word/_rels/document.xml.rels"] = parts["word/_rels/document.xml.rels"].replace(
        b"</Relationships>",
        b'<Relationship Id="rIdEn" Type="http://schemas.openxmlformats.org/officeDocument/'
        b'2006/relationships/endnotes" Target="endnotes.xml"/></Relationships>')
    parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(
        b"</Types>",
        b'<Override PartName="/word/endnotes.xml" ContentType="application/vnd.openxmlformats-'
        b'officedocument.wordprocessingml.endnotes+xml"/></Types>')
    dst = tmp_path / "with-endnote.docx"
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in parts.items():
            zout.writestr(name, data)

    content = docx_to_ast(dst)["body"][0]["content"]

    assert [n["type"] for n in content] == ["paragraph", "endnote"]
    assert content[1]["content"][0]["text"] == "An endnote's words."
