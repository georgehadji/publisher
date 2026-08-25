"""
Rich DOCX extraction: the things `para.text` threw away.

Every case here failed before `docx_rich` existed, and failed *silently*: the
old extractor's no-loss post-condition compares prose characters, and none of
bold, a table's structure, an embedded picture or a footnote body is a prose
character in the body paragraph it belongs to.
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import docx
import pytest

from publisher_ingest.docx_to_ast import IngestError, docx_to_ast, read_blocks
from publisher_ingest.docx_rich import sha256_hex

# 1x1 transparent PNG. Small enough to inline, real enough for python-docx to
# read an image header off.
PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08"
    b"\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00"
    b"\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
)

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

FOOTNOTES_XML = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:footnotes xmlns:w="{W_NS}">
  <w:footnote w:type="separator" w:id="0"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>
  <w:footnote w:id="2"><w:p><w:r><w:t>Kant disagrees, sharply.</w:t></w:r></w:p></w:footnote>
</w:footnotes>""".encode("utf-8")


def _manuscript(tmp_path: Path, *, with_image: bool = False) -> Path:
    """A DOCX shaped like the corpus: heading, prose, and rich content.

    The prose paragraph clears SUBSTANTIVE_PARAGRAPH_CHARS so the section reads
    as body rather than front matter.
    """
    document = docx.Document()
    document.add_paragraph("CHAPTER ONE")

    para = document.add_paragraph()
    para.add_run("Plain and ")
    para.add_run("italic").italic = True
    para.add_run(" and ")
    para.add_run("bold").bold = True
    para.add_run(". " + "Filler prose to clear the substantive threshold. " * 8)

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Copies"
    table.cell(1, 0).text = "Attica"
    table.cell(1, 1).text = "1200"

    if with_image:
        image = tmp_path / "plate.png"
        image.write_bytes(PNG)
        document.add_picture(str(image))

    path = tmp_path / "manuscript.docx"
    document.save(str(path))
    return path


def _add_footnote(src: Path, dst: Path, *, anchor: str = "bold") -> Path:
    """Attach a real footnote to the run containing `anchor`.

    python-docx cannot author footnotes, so the part, its relationship, its
    content-type override and the in-body reference are written directly. This
    is also the only way to get a fixture whose footnote text lives outside
    `document.xml` -- the case the extractor has to handle.
    """
    with zipfile.ZipFile(src) as zin:
        parts = {n: zin.read(n) for n in zin.namelist()}

    body = parts["word/document.xml"].decode("utf-8")
    close = body.index("</w:r>", body.index(f">{anchor}<")) + len("</w:r>")
    body = (
        body[:close]
        + '<w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr>'
          '<w:footnoteReference w:id="2"/></w:r>'
        + body[close:]
    )
    parts["word/document.xml"] = body.encode("utf-8")
    parts["word/footnotes.xml"] = FOOTNOTES_XML

    rels = parts["word/_rels/document.xml.rels"].decode("utf-8")
    parts["word/_rels/document.xml.rels"] = rels.replace(
        "</Relationships>",
        '<Relationship Id="rIdFn" Type="http://schemas.openxmlformats.org/'
        'officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/>'
        "</Relationships>",
    ).encode("utf-8")

    types = parts["[Content_Types].xml"].decode("utf-8")
    parts["[Content_Types].xml"] = types.replace(
        "</Types>",
        '<Override PartName="/word/footnotes.xml" ContentType="application/'
        'vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>'
        "</Types>",
    ).encode("utf-8")

    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in parts.items():
            zout.writestr(name, data)
    return dst


def _walk(node, out: list) -> None:
    if isinstance(node, list):
        for n in node:
            _walk(n, out)
    elif isinstance(node, dict):
        out.append(node)
        for key in ("frontMatter", "body", "backMatter", "content"):
            if node.get(key):
                _walk(node[key], out)


def _nodes_of(root, ntype: str) -> list[dict]:
    found: list[dict] = []
    _walk(root, found)
    return [n for n in found if n.get("type") == ntype]


def _all_text(root) -> str:
    return "".join(t.get("text", "") for t in _nodes_of(root, "text"))


def test_runs_keep_their_marks(tmp_path):
    """`para.text` concatenates runs and drops `w:rPr`, so before this every
    italic and bold in every manuscript was flattened to plain prose."""
    ast = docx_to_ast(_manuscript(tmp_path))
    texts = _nodes_of(ast, "text")
    marked = {t["text"]: [m["type"] for m in t["marks"]] for t in texts if t.get("marks")}
    assert marked.get("italic") == ["emphasis"]
    assert marked.get("bold") == ["strong"]
    # Adjacent unmarked runs coalesce instead of becoming one node per run.
    assert any(t["text"].startswith("Plain and ") for t in texts)


def test_tables_survive_as_tables(tmp_path):
    """Cells used to be unrolled into loose paragraphs: the text survived, the
    table did not, and no gate noticed because the prose was all present."""
    ast = docx_to_ast(_manuscript(tmp_path))
    tables = _nodes_of(ast, "table")
    assert len(tables) == 1
    rows = tables[0]["content"]
    assert [r["type"] for r in rows] == ["table-row", "table-row"]
    assert len(rows[0]["content"]) == 2
    assert {t["text"] for t in _nodes_of(rows[1], "text")} == {"Attica", "1200"}


def test_images_are_stored_and_referenced(tmp_path):
    """A picture carries no text, so dropping one tripped no check at all."""
    stored: dict[str, bytes] = {}

    def sink(data: bytes, media_type: str, name: str) -> str:
        digest = sha256_hex(data)
        stored[digest] = data
        return digest

    ast = docx_to_ast(_manuscript(tmp_path, with_image=True), store_media=sink)
    figures = _nodes_of(ast, "figure")
    assert len(figures) == 1
    ref = figures[0]["attrs"]["mediaRef"]
    assert ref["mediaType"] == "image/png"
    assert re.fullmatch(r"[0-9a-f]{64}", ref["hash"])
    assert stored[ref["hash"]] == PNG


def test_an_image_with_nowhere_to_go_is_an_error_not_a_shrug(tmp_path):
    """Refusing beats dropping: a manuscript that silently loses its plates is
    the failure this extractor exists to end."""
    with pytest.raises(IngestError, match="media sink"):
        docx_to_ast(_manuscript(tmp_path, with_image=True))


def test_footnotes_are_lifted_out_of_their_own_part(tmp_path):
    """Footnote bodies live in `word/footnotes.xml`, which python-docx does not
    expose -- so their text never reached the AST, unseen by every check."""
    ast = docx_to_ast(_add_footnote(_manuscript(tmp_path), tmp_path / "note.docx"))

    notes = _nodes_of(ast, "footnote")
    assert len(notes) == 1
    assert notes[0]["attrs"]["number"] == 1
    assert _all_text(notes[0]) == "Kant disagrees, sharply."

    # No marker text spliced into the prose: every target engine numbers its
    # own call, so a literal "1" after "bold" would print twice.
    prose = [p for p in _nodes_of(ast, "paragraph") if "Plain and" in _all_text(p)]
    assert prose and "bold. " in _all_text(prose[0])
    assert "bold1" not in _all_text(prose[0])


def test_a_footnote_on_a_chapter_heading_is_not_dropped(tmp_path):
    """A heading contributes its text as the section title and its paragraph
    node is discarded -- which used to take any footnote or figure hanging off
    the same `w:p` down with it. The note opens the section instead."""
    src = _add_footnote(_manuscript(tmp_path), tmp_path / "h.docx",
                        anchor="CHAPTER ONE")
    ast = docx_to_ast(src)

    assert _all_text(_nodes_of(ast, "footnote")) == "Kant disagrees, sharply."
    # The synthesised marker must not contaminate the chapter title.
    titles = [c["attrs"]["title"] for c in _nodes_of(ast, "chapter")]
    assert titles == ["CHAPTER ONE"]


def test_source_text_still_cannot_be_lost(tmp_path):
    """The no-loss post-condition holds over the richer output: table cells and
    footnote bodies are now part of what has to survive."""
    src = _add_footnote(_manuscript(tmp_path), tmp_path / "n.docx")
    _, sources = read_blocks(src)
    ast = docx_to_ast(src)

    # A heading's text becomes `attrs.title`, not a text node, so the haystack
    # has to include titles -- exactly as the production check does.
    emitted = _all_text(ast) + "\n".join(
        n["attrs"]["title"] for n in _nodes_of(ast, "chapter")
    )
    assert "Attica" in sources
    for text in sources:
        assert text in emitted
